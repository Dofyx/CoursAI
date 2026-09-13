import sys,json,re,shutil,subprocess,threading,webbrowser,mimetypes
from pathlib import Path
from datetime import datetime
import requests,sounddevice as sd,soundfile as sf
from docx import Document
from docx.shared import Pt
from PySide6.QtCore import Signal,QObject,Qt
from PySide6.QtGui import QIcon,QPixmap
from PySide6.QtWidgets import *
try: import keyring
except: keyring=None
APP='CoursIA'; BASE=Path.home()/'Documents'/'CoursIA'; REC=BASE/'Enregistrements'; TR=BASE/'Transcriptions'; DOC=BASE/'Documents'; META=BASE/'Metadonnees'; CFG=BASE/'config.json'
for p in (REC,TR,DOC,META):p.mkdir(parents=True,exist_ok=True)
DEF={'subjects':['Droit','Économie','Histoire','Informatique','Mathématiques'],'mistral_transcription_model':'voxtral-mini-latest','mistral_text_model':'mistral-small-latest'}
def cfgload():
 try:return {**DEF,**json.loads(CFG.read_text(encoding='utf-8'))}
 except:return DEF.copy()
def cfgsave(c):CFG.write_text(json.dumps(c,ensure_ascii=False,indent=2),encoding='utf-8')
def secret(n,c):
 if keyring:
  try:return keyring.get_password(APP,n) or ''
  except:pass
 return c.get(n,'')
def putsecret(n,v,c):
 if keyring:
  try:keyring.set_password(APP,n,v);c.pop(n,None);return
  except:pass
 c[n]=v
def clean(s):return re.sub(r'[\\/:*?"<>|]+','_',s.strip()).strip(' .') or 'Sans titre'
def stem(subject,title,dt=None):return f'{clean(subject)} - {clean(title)} - {(dt or datetime.now()).strftime("%Y-%m-%d_%H-%M")}'
def meta_path(st):return META/(st+'.json')
def write_meta(st,subject,title,source):meta_path(st).write_text(json.dumps({'matiere':subject,'titre':title,'date_heure':datetime.now().isoformat(timespec='minutes'),'source':source},ensure_ascii=False,indent=2),encoding='utf-8')
def read_meta(st):
 try:return json.loads(meta_path(st).read_text(encoding='utf-8'))
 except:return {'matiere':'Non renseignée','titre':st,'date_heure':'Non renseignée','source':'inconnu'}
def xopen(p):subprocess.Popen(['xdg-open',str(p)])
def asset(name):return Path(__file__).resolve().parent/'assets'/name
class Bus(QObject):done=Signal(object);error=Signal(str);status=Signal(str)
class Recorder:
 def __init__(self):self.stream=self.file=None;self.paused=False
 def start(self,p,d):
  self.file=sf.SoundFile(p,'w',samplerate=48000,channels=1,subtype='PCM_16')
  def cb(x,n,t,s):
   if not self.paused:self.file.write(x.copy())
  self.stream=sd.InputStream(device=d,channels=1,samplerate=48000,callback=cb);self.stream.start()
 def pause(self):self.paused=True
 def resume(self):self.paused=False
 def stop(self):
  if self.stream:self.stream.stop();self.stream.close();self.stream=None
  if self.file:self.file.close();self.file=None
class Manager(QGroupBox):
 def __init__(self,title,folder,pats):
  super().__init__(title);self.folder=folder;self.pats=pats;l=QVBoxLayout(self);self.list=QListWidget();l.addWidget(self.list);h=QHBoxLayout()
  for t,f in [('Ouvrir',self.open),('Renommer',self.rename),('Supprimer',self.delete)]:b=QPushButton(t);b.clicked.connect(f);h.addWidget(b)
  l.addLayout(h);self.refresh()
 def paths(self):
  a=[]
  for x in self.pats:a+=list(self.folder.glob(x))
  return sorted(set(a),key=lambda p:p.stat().st_mtime,reverse=True)
 def refresh(self):self.list.clear();self.list.addItems([p.name for p in self.paths()])
 def selected(self):return self.folder/self.list.currentItem().text() if self.list.currentItem() else None
 def open(self):
  if self.selected():xopen(self.selected())
 def rename(self):
  p=self.selected()
  if not p:return
  n,ok=QInputDialog.getText(self,'Renommer','Nouveau nom',text=p.stem)
  if ok and n:
   old=p.stem;p.rename(p.with_name(clean(n)+p.suffix));m=meta_path(old)
   if m.exists():m.rename(meta_path(clean(n)))
   self.refresh()
 def delete(self):
  p=self.selected()
  if p and QMessageBox.question(self,'Suppression',f'Supprimer {p.name} ?')==QMessageBox.Yes:
   m=meta_path(p.stem);p.unlink();m.unlink(missing_ok=True);self.refresh()
class Main(QMainWindow):
 def __init__(self):
  super().__init__();self.c=cfgload();self.r=Recorder();self.bus=Bus();self.bus.done.connect(self.finished);self.bus.error.connect(self.err);self.bus.status.connect(self.statusBar().showMessage);self.resize(1200,780);self.setWindowTitle('CoursIA 0.5')
  app_icon=asset('logo_icon.ico') if asset('logo_icon.ico').exists() else asset('logo_icon_256.png')
  if app_icon.exists():self.setWindowIcon(QIcon(str(app_icon)))
  root=QWidget();self.setCentralWidget(root);h=QHBoxLayout(root);h.setContentsMargins(0,0,0,0);side=QFrame();side.setObjectName('side');side.setFixedWidth(250);sl=QVBoxLayout(side);side_logo=QLabel();side_logo.setAlignment(Qt.AlignCenter);pix=QPixmap(str(asset('logo_icon_128.png')))
  if not pix.isNull():side_logo.setPixmap(pix.scaled(88,88,Qt.KeepAspectRatio,Qt.SmoothTransformation))
  sl.addWidget(side_logo);brand=QLabel('CoursIA');brand.setObjectName('brand');brand.setAlignment(Qt.AlignCenter);sl.addWidget(brand);subtitle=QLabel('Assistant de cours');subtitle.setAlignment(Qt.AlignCenter);sl.addWidget(subtitle);self.stack=QStackedWidget();pages=[('Enregistrement',self.pg_rec()),('Retranscriptions',self.pg_tr()),('Résumés & révisions',self.pg_doc()),('Paramètres',self.pg_cfg()),('À propos',self.pg_about())]
  for i,(name,page) in enumerate(pages):b=QPushButton(name);b.setCheckable(True);b.setAutoExclusive(True);b.clicked.connect(lambda _,z=i:self.stack.setCurrentIndex(z));sl.addWidget(b);self.stack.addWidget(page);b.setChecked(i==0)
  sl.addStretch();h.addWidget(side);h.addWidget(self.stack,1);self.statusBar().showMessage('Prêt')
 def head(self,t,s):w=QWidget();l=QVBoxLayout(w);a=QLabel(t);a.setObjectName('title');l.addWidget(a);l.addWidget(QLabel(s));return w
 def confirm_overwrite(self,path,label):
  if not path.exists():return True
  return QMessageBox.warning(self,'Fichier existant',f'{label} existe déjà :\n\n{path.name}\n\nVoulez-vous l’écraser ? Cette action est irréversible.',QMessageBox.Yes|QMessageBox.No,QMessageBox.No)==QMessageBox.Yes
 def pg_rec(self):
  w=QWidget();l=QVBoxLayout(w);l.addWidget(self.head('Gestion de l’enregistrement','Nommage automatique : matière, titre, date et heure.'));g=QGroupBox('Nouveau cours');f=QFormLayout(g);self.subject=QComboBox();f.addRow('Matière',self.subject);self.course=QLineEdit();self.course.setPlaceholderText('Ex. Introduction au droit constitutionnel');self.course.textChanged.connect(self.preview);self.subject.currentTextChanged.connect(self.preview);f.addRow('Titre du cours',self.course);self.name_preview=QLabel();self.name_preview.setWordWrap(True);f.addRow('Nom produit',self.name_preview);self.mic=QComboBox()
  for i,d in enumerate(sd.query_devices()):
   if d['max_input_channels']>0:self.mic.addItem(f"{d['name']} ({i})",i)
  f.addRow('Microphone',self.mic);h=QHBoxLayout()
  for t,fn in [('● Commencer',self.start),('Pause',lambda:(self.r.pause(),self.statusBar().showMessage('En pause'))),('Reprendre',lambda:(self.r.resume(),self.statusBar().showMessage('Repris'))),('Fin',self.stop),('Ajouter un fichier audio',self.import_audio)]:b=QPushButton(t);b.clicked.connect(fn);h.addWidget(b)
  f.addRow(h);l.addWidget(g);self.fmrec=Manager('Enregistrements',REC,['*.wav','*.mp3','*.m4a','*.flac','*.ogg','*.opus']);l.addWidget(self.fmrec,1);self.loadsubjects();return w
 def loadsubjects(self):
  cur=self.subject.currentText() if hasattr(self,'subject') else '';self.subject.clear();self.subject.addItems(self.c.get('subjects',[]));i=self.subject.findText(cur);self.subject.setCurrentIndex(max(0,i));self.preview()
 def preview(self):
  if hasattr(self,'name_preview'):self.name_preview.setText(stem(self.subject.currentText() or 'Sans matière',self.course.text() or 'Titre du cours')+'.wav')
 def checked_identity(self):
  if not self.subject.currentText():raise ValueError('Ajoutez et sélectionnez une matière dans Paramètres.')
  if not self.course.text().strip():raise ValueError('Renseignez le titre du cours.')
  return self.subject.currentText(),self.course.text().strip(),stem(self.subject.currentText(),self.course.text())
 def start(self):
  try:s,t,st=self.checked_identity();self.r.start(REC/(st+'.wav'),self.mic.currentData());write_meta(st,s,t,'enregistrement');self.statusBar().showMessage('Enregistrement en cours…')
  except Exception as e:self.err(str(e))
 def stop(self):self.r.stop();self.fmrec.refresh();self.refresh();self.statusBar().showMessage('Enregistrement terminé')
 def import_audio(self):
  try:
   s,t,st=self.checked_identity();src,_=QFileDialog.getOpenFileName(self,'Ajouter un fichier audio','','Audio (*.wav *.mp3 *.m4a *.flac *.ogg *.opus);;Tous les fichiers (*)')
   if not src:return
   ext=Path(src).suffix.lower();dst=REC/(st+ext)
   if not self.confirm_overwrite(dst,'Le fichier audio'):return
   shutil.copy2(src,dst);write_meta(st,s,t,'fichier importé');self.fmrec.refresh();self.refresh();QMessageBox.information(self,'Fichier ajouté',f'Le fichier a été copié et renommé :\n{dst.name}')
  except Exception as e:self.err(str(e))
 def pg_tr(self):
  w=QWidget();l=QVBoxLayout(w);l.addWidget(self.head('Gestion des retranscriptions','Transcription directe avec Mistral Voxtral.'));g=QGroupBox('Retranscrire');f=QFormLayout(g);f.addRow('IA',QLabel('Mistral Voxtral'));self.trrec=QComboBox();f.addRow('Enregistrement',self.trrec);b=QPushButton('Retranscrire');b.clicked.connect(self.transcribe);f.addRow('',b);l.addWidget(g);self.fmtr=Manager('Retranscriptions',TR,['*.txt']);l.addWidget(self.fmtr,1);return w
 def transcribe(self):
  self.refresh();n=self.trrec.currentText();k=secret('mistral_api_key',self.c)
  if not n:return self.err('Aucun enregistrement sélectionné.')
  if not k:return self.err('Renseignez la clé Mistral dans Paramètres.')
  source=REC/n;out=TR/(source.stem+'.txt')
  if not self.confirm_overwrite(out,'La retranscription'):return
  self.bg(self.do_transcribe,source,k)
 def mistral_chat(self,k,prompt,timeout=1800):
  model=self.c.get('mistral_text_model','mistral-small-latest').strip() or 'mistral-small-latest';r=requests.post('https://api.mistral.ai/v1/chat/completions',headers={'Authorization':f'Bearer {k}','Content-Type':'application/json'},json={'model':model,'messages':[{'role':'user','content':prompt}],'temperature':0.2},timeout=timeout)
  if not r.ok:
   try:detail=r.json().get('message') or r.json().get('error',{}).get('message') or r.text
   except Exception:detail=r.text
   raise RuntimeError(f'Erreur Mistral {r.status_code} : {detail}')
  return r.json()['choices'][0]['message']['content']
 def do_transcribe(self,p,k):
  self.bus.status.emit('Transcription Mistral en cours…');model=self.c.get('mistral_transcription_model','voxtral-mini-latest').strip() or 'voxtral-mini-latest';mime=mimetypes.guess_type(p.name)[0] or 'application/octet-stream'
  with open(p,'rb') as f:r=requests.post('https://api.mistral.ai/v1/audio/transcriptions',headers={'Authorization':f'Bearer {k}'},files={'file':(p.name,f,mime)},data={'model':model,'language':'fr'},timeout=7200)
  if not r.ok:
   try:detail=r.json().get('message') or r.json().get('error',{}).get('message') or r.text
   except Exception:detail=r.text
   raise RuntimeError(f'Erreur Mistral {r.status_code} : {detail}')
  out=TR/(p.stem+'.txt');out.write_text(r.json()['text'],encoding='utf-8');self.bus.done.emit(('tr',out))
 def pg_doc(self):
  w=QWidget();l=QVBoxLayout(w);l.addWidget(self.head('Résumés et fiches de révision','Deux productions pédagogiques distinctes avec Mistral.'));g=QGroupBox('Générer');f=QFormLayout(g);f.addRow('IA',QLabel('Mistral'));self.doctr=QComboBox();f.addRow('Retranscription',self.doctr);h=QHBoxLayout()
  for text,kind in [('Créer le résumé','resume'),('Créer la fiche de révision','fiche')]:b=QPushButton(text);b.clicked.connect(lambda _,x=kind:self.document(x));h.addWidget(b)
  f.addRow(h);l.addWidget(g);self.fmdoc=Manager('Documents Word',DOC,['*.docx']);l.addWidget(self.fmdoc,1);return w
 def document(self,kind):
  self.refresh();n=self.doctr.currentText();k=secret('mistral_api_key',self.c)
  if not n:return self.err('Aucune retranscription sélectionnée.')
  if not k:return self.err('Renseignez la clé Mistral dans Paramètres.')
  source=TR/n;suffix='_resume.docx' if kind=='resume' else '_fiche_revision.docx';out=DOC/(source.stem+suffix)
  if not self.confirm_overwrite(out,'Le document Word'):return
  self.bg(self.do_document,source,kind,k)
 def do_document(self,p,kind,k):
  m=read_meta(p.stem);text=p.read_text(encoding='utf-8');context=f"Matière : {m['matiere']}\nTitre : {m['titre']}\nDate : {m['date_heure']}"
  if kind=='resume':prompt=f"""{context}\n\nProduis un résumé universitaire rédigé et cohérent. Structure imposée : 1. Objet et problématique, 2. Plan suivi, 3. Synthèse développée des idées et raisonnements, 4. Exemples pédagogiques utiles, 5. Conclusion et points essentiels. Rédige en paragraphes complets, conserve les nuances et n'invente rien. N'ajoute ni quiz ni cartes mémoire.\n\nTRANSCRIPTION :\n{text}""";title='Résumé du cours'
  else:prompt=f"""{context}\n\nCrée une fiche de révision opérationnelle, concise et mémorisable, non narrative. Structure imposée : 1. Objectifs à maîtriser, 2. Définitions essentielles, 3. Concepts et mécanismes en listes, 4. Repères, auteurs, dates ou formules présents, 5. Exemples à retenir, 6. Erreurs à éviter, 7. Questions-réponses, 8. Mini-quiz corrigé, 9. Checklist « Je sais… ». N'invente rien.\n\nTRANSCRIPTION :\n{text}""";title='Fiche de révision'
  self.bus.status.emit(f'Génération de : {title}…');content=self.mistral_chat(k,prompt);out=DOC/(p.stem+('_resume.docx' if kind=='resume' else '_fiche_revision.docx'));d=Document();d.add_heading(f"{title} - {m['matiere']}",0);d.add_paragraph(m['titre']);d.add_paragraph(f"Date du cours : {m['date_heure']}")
  for line in content.splitlines():
   x=line.strip()
   if not x:continue
   if x.startswith('### '):d.add_heading(x[4:],3)
   elif x.startswith('## '):d.add_heading(x[3:],2)
   elif x.startswith('# '):d.add_heading(x[2:],1)
   elif re.match(r'^[-*] ',x):d.add_paragraph(x[2:],style='List Bullet')
   elif re.match(r'^\d+[.)] ',x):d.add_paragraph(re.sub(r'^\d+[.)] ','',x),style='List Number')
   else:d.add_paragraph(x.replace('**',''))
  d.styles['Normal'].font.name='Aptos';d.styles['Normal'].font.size=Pt(11);d.save(out);self.bus.done.emit(('doc',out))
 def pg_cfg(self):
  page=QWidget();outer=QVBoxLayout(page);scroll=QScrollArea();scroll.setWidgetResizable(True);body=QWidget();l=QVBoxLayout(body);l.addWidget(self.head('Paramètres','Matières et accès Mistral Pay As You Go.'));g=QGroupBox('Matières, une par ligne');gl=QVBoxLayout(g);self.subjects=QPlainTextEdit('\n'.join(self.c.get('subjects',[])));self.subjects.setMinimumHeight(150);gl.addWidget(self.subjects);l.addWidget(g);g=QGroupBox('Mistral AI');f=QFormLayout(g);self.mistral_key=QLineEdit(secret('mistral_api_key',self.c));self.mistral_key.setEchoMode(QLineEdit.Password);f.addRow('Clé API',self.mistral_key);b=QPushButton('Ouvrir la console Mistral');b.clicked.connect(lambda:webbrowser.open('https://console.mistral.ai/'));f.addRow('',b);l.addWidget(g);g=QGroupBox('Modèles Mistral');f=QFormLayout(g);self.mtm=QLineEdit(self.c.get('mistral_transcription_model','voxtral-mini-latest'));f.addRow('Modèle de transcription',self.mtm);self.mxm=QLineEdit(self.c.get('mistral_text_model','mistral-small-latest'));f.addRow('Modèle de résumé',self.mxm);l.addWidget(g);b=QPushButton('Enregistrer les paramètres');b.clicked.connect(self.save);l.addWidget(b);l.addStretch();scroll.setWidget(body);outer.addWidget(scroll);return page
 def save(self):
  subjects=[]
  for x in self.subjects.toPlainText().splitlines():
   x=x.strip()
   if x and x not in subjects:subjects.append(x)
  if not subjects:return self.err('Ajoutez au moins une matière.')
  self.c={'subjects':subjects,'mistral_transcription_model':self.mtm.text().strip() or 'voxtral-mini-latest','mistral_text_model':self.mxm.text().strip() or 'mistral-small-latest'};putsecret('mistral_api_key',self.mistral_key.text().strip(),self.c);cfgsave(self.c);self.loadsubjects();self.statusBar().showMessage('Paramètres enregistrés')
 def pg_about(self):
  w=QWidget();l=QVBoxLayout(w);l.setContentsMargins(70,35,70,35);l.addStretch();logo=QLabel();logo.setAlignment(Qt.AlignCenter);pix=QPixmap(str(asset('logo.png')))
  if not pix.isNull():logo.setPixmap(pix.scaled(220,220,Qt.KeepAspectRatio,Qt.SmoothTransformation))
  l.addWidget(logo);brand=QLabel('CoursIA');brand.setAlignment(Qt.AlignCenter);brand.setStyleSheet('font-size:34px;font-weight:800;color:#62d0ff;');l.addWidget(brand);version=QLabel('Version 0.5');version.setAlignment(Qt.AlignCenter);version.setStyleSheet('font-size:18px;color:#9fb3cc;margin-bottom:18px;');l.addWidget(version);tag=QLabel('Enregistrement, transcription et révision universitaire assistés par Mistral AI');tag.setWordWrap(True);tag.setAlignment(Qt.AlignCenter);l.addWidget(tag);card=QGroupBox('Licence');cl=QVBoxLayout(card);lic=QLabel('<a href="https://creativecommons.org/licenses/by-nc-sa/4.0/deed.fr" style="color:#62d0ff;">Creative Commons Attribution - Pas d’Utilisation Commerciale - Partage dans les Mêmes Conditions 4.0 International</a>');lic.setTextFormat(Qt.RichText);lic.setTextInteractionFlags(Qt.TextBrowserInteraction);lic.setOpenExternalLinks(True);lic.setWordWrap(True);lic.setAlignment(Qt.AlignCenter);cl.addWidget(lic);note=QLabel('CC BY-NC-SA 4.0');note.setAlignment(Qt.AlignCenter);cl.addWidget(note);l.addWidget(card);credit=QLabel('Developed by Dofyx AI Corp');credit.setAlignment(Qt.AlignCenter);credit.setStyleSheet('font-size:20px;font-weight:700;color:white;margin-top:22px;');l.addWidget(credit);github=QLabel('<a href="https://github.com/Dofyx/CoursAI" style="color:#62d0ff;">GitHub - Dofyx/CoursAI</a>');github.setTextFormat(Qt.RichText);github.setTextInteractionFlags(Qt.TextBrowserInteraction);github.setOpenExternalLinks(True);github.setAlignment(Qt.AlignCenter);github.setStyleSheet('font-size:15px;margin-top:10px;');l.addWidget(github);l.addStretch();return w
 def refresh(self):
  self.trrec.clear();self.trrec.addItems([p.name for p in self.fmrec.paths()]);self.doctr.clear();self.doctr.addItems([p.name for p in self.fmtr.paths()])
 def rate(self,r,p):
  a=[]
  for h,n in [('x-ratelimit-remaining-requests','requêtes'),('x-ratelimit-remaining-tokens','jetons')]:
   if r.headers.get(h):a.append(f"{r.headers[h]} {n} restants")
  self.quota[p]=' · '.join(a) or 'non communiqué sur cette réponse'
 def showquota(self):
  if hasattr(self,'trquota'):self.trquota.setText('Quota : '+self.quota.get('groq','aucune mesure récente'))
  if hasattr(self,'docquota'):self.docquota.setText('Quota : '+self.quota.get(self.docai.currentData(),'aucune mesure récente'))
 def bg(self,f,*a):threading.Thread(target=self.worker,args=(f,a),daemon=True).start()
 def worker(self,f,a):
  try:f(*a)
  except Exception as e:self.bus.error.emit(str(e))
 def finished(self,x):
  k,p=x
  if k=='tr':self.fmtr.refresh()
  else:self.fmdoc.refresh()
  self.refresh();self.statusBar().showMessage('Terminé : '+p.name)
 def err(self,s):QMessageBox.critical(self,'Erreur',s)
 def showEvent(self,e):super().showEvent(e);self.refresh()
 def closeEvent(self,e):self.r.stop();e.accept()
STYLE="""QWidget{background:#0b1220;color:#e6edf7;font-family:Inter,Segoe UI,sans-serif;font-size:14px}#side{background:#111a2b;border-right:1px solid #26344d}#brand{font-size:28px;font-weight:700;color:#62d0ff;margin:20px 8px 2px}#side QPushButton{padding:14px;text-align:left;border:0;border-radius:9px;margin:3px;background:transparent}#side QPushButton:hover{background:#1a2942}#side QPushButton:checked{background:#1677ff;color:white}QLabel#title{font-size:26px;font-weight:700;color:white;margin-top:12px}QGroupBox{border:1px solid #26344d;border-radius:12px;margin-top:16px;padding:16px;background:#101a2a;font-weight:600}QGroupBox::title{subcontrol-origin:margin;left:14px;padding:0 6px;color:#8bdcff}QPushButton{background:#1677ff;border:0;border-radius:8px;padding:9px 14px;font-weight:600}QPushButton:hover{background:#4096ff}QLineEdit,QComboBox,QListWidget,QPlainTextEdit{background:#0b1424;border:1px solid #31415f;border-radius:8px;padding:8px;selection-background-color:#1677ff}QStatusBar{background:#101a2a;color:#9fb3cc}"""
if __name__=='__main__':a=QApplication(sys.argv);a.setStyleSheet(STYLE);m=Main();m.show();sys.exit(a.exec())
