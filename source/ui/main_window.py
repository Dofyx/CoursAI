import sys
import threading
import webbrowser
import re
import mimetypes
from pathlib import Path
from datetime import datetime

from PySide6.QtCore import Signal, QObject, Qt, QSize
from PySide6.QtGui import QIcon, QPixmap
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, 
    QStackedWidget, QFrame, QLabel, QPushButton, QMessageBox, 
    QStatusBar, QProgressBar
)

from config.settings import load_config, save_config, get_secret, set_secret
from config.paths import BASE, REC, TR, DOC, META, CFG, APP
from audio.recorder import Recorder
from audio.manager import Manager
from mistral.client import MistralClient
from mistral.prompts import load_prompt, render_prompt
from utils.helpers import open_file, clean, stem
from utils.meta import write_meta, read_meta
from .styles import STYLE


class Bus(QObject):
    """Système de signaux pour la communication entre threads."""
    done = Signal(object)
    error = Signal(str)
    status = Signal(str)


class MainWindow(QMainWindow):
    """Fenêtre principale de l'application."""
    
    def __init__(self):
        super().__init__()
        self.c = load_config()
        self.r = Recorder()
        self.bus = Bus()
        self.bus.done.connect(self._finished)
        self.bus.error.connect(self._err)
        self.bus.status.connect(self.statusBar().showMessage)
        
        # Barre de progression
        self.progress_bar = QProgressBar()
        self.progress_bar.setVisible(False)
        self.statusBar().addPermanentWidget(self.progress_bar)
        
        # Initialisation de l'UI
        self._init_ui()

    def _init_ui(self):
        """Initialise l'interface utilisateur."""
        self.resize(1200, 780)
        self.setWindowTitle('CoursIA 0.6')
        
        # Icône de l'application
        app_icon = self._asset('logo_icon.ico')
        if not app_icon.exists():
            app_icon = self._asset('logo_icon_256.png')
        if app_icon.exists():
            self.setWindowIcon(QIcon(str(app_icon)))
        
        # Widget central
        root = QWidget()
        self.setCentralWidget(root)
        h = QHBoxLayout(root)
        h.setContentsMargins(0, 0, 0, 0)
        
        # Barre latérale
        side = QFrame()
        side.setObjectName('side')
        side.setFixedWidth(250)
        sl = QVBoxLayout(side)
        
        # Logo
        side_logo = QLabel()
        side_logo.setAlignment(Qt.AlignCenter)
        pix = QPixmap(str(self._asset('logo_icon_128.png')))
        if not pix.isNull():
            side_logo.setPixmap(pix.scaled(88, 88, Qt.KeepAspectRatio, Qt.SmoothTransformation))
        sl.addWidget(side_logo)
        
        # Titre
        brand = QLabel('CoursIA')
        brand.setObjectName('brand')
        brand.setAlignment(Qt.AlignCenter)
        sl.addWidget(brand)
        
        # Sous-titre
        subtitle = QLabel('Assistant de cours')
        subtitle.setAlignment(Qt.AlignCenter)
        sl.addWidget(subtitle)
        
        # Pages
        self.stack = QStackedWidget()
        self.rec_page = None
        self.tr_page = None
        self.doc_page = None
        self.cfg_page = None
        self.about_page = None
        
        pages = [
            ('Enregistrement', self._create_rec_page()),
            ('Retranscriptions', self._create_tr_page()),
            ('Résumés & révisions', self._create_doc_page()),
            ('Paramètres', self._create_cfg_page()),
            ('À propos', self._create_about_page())
        ]
        
        for i, (name, page) in enumerate(pages):
            b = QPushButton(name)
            b.setCheckable(True)
            b.setAutoExclusive(True)
            b.clicked.connect(lambda _, z=i: self.stack.setCurrentIndex(z))
            sl.addWidget(b)
            self.stack.addWidget(page)
            b.setChecked(i == 0)
        
        sl.addStretch()
        h.addWidget(side)
        h.addWidget(self.stack, 1)
        
        self.statusBar().showMessage('Prêt')

    def _asset(self, name: str) -> Path:
        """Retourne le chemin d'un asset."""
        return Path(__file__).parent.parent / 'assets' / name

    def _create_rec_page(self):
        """Crée la page d'enregistrement."""
        from .pages.rec import RecPage
        self.rec_page = RecPage(self, self.bus, self.c, self.r)
        return self.rec_page

    def _create_tr_page(self):
        """Crée la page de retranscription."""
        from .pages.tr import TrPage
        self.tr_page = TrPage(self, self.bus, self.c)
        return self.tr_page

    def _create_doc_page(self):
        """Crée la page de génération de documents."""
        from .pages.doc import DocPage
        self.doc_page = DocPage(self, self.bus, self.c)
        return self.doc_page

    def _create_cfg_page(self):
        """Crée la page de configuration."""
        from .pages.cfg import CfgPage
        self.cfg_page = CfgPage(self, self.c)
        return self.cfg_page

    def _create_about_page(self):
        """Crée la page "À propos"."""
        from .pages.about import AboutPage
        self.about_page = AboutPage(self)
        return self.about_page

    def head(self, t: str, s: str) -> QWidget:
        """Crée un en-tête de page."""
        w = QWidget()
        l = QVBoxLayout(w)
        a = QLabel(t)
        a.setObjectName('title')
        l.addWidget(a)
        l.addWidget(QLabel(s))
        return w

    def confirm_overwrite(self, path: Path, label: str) -> bool:
        """Demande confirmation avant écrasement."""
        if not path.exists():
            return True
        return QMessageBox.warning(
            self, 'Fichier existant',
            f'{label} existe déjà :\n\n{path.name}\n\nVoulez-vous l’écraser ? Cette action est irréversible.',
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.No
        ) == QMessageBox.Yes

    def _do_transcribe(self, source: Path, k: str):
        """Effectue la transcription via Mistral."""
        self.bus.status.emit('Transcription Mistral en cours…')
        model = self.c.get('mistral_transcription_model', 'voxtral-mini-latest').strip() or 'voxtral-mini-latest'
        mime = mimetypes.guess_type(source.name)[0] or 'application/octet-stream'
        
        client = MistralClient(k)
        try:
            text = client.transcribe(source, model=model, language='fr')
            out = TR / (source.stem + '.txt')
            out.write_text(text, encoding='utf-8')
            self.bus.done.emit(('tr', out))
        except Exception as e:
            self.bus.error.emit(str(e))

    def _do_document(self, p: Path, kind: str, k: str):
        """Génère un document (résumé ou fiche)."""
        m = read_meta(p.stem)
        text = p.read_text(encoding='utf-8')
        context = f"Matière : {m['matiere']}\nTitre : {m['titre']}\nDate : {m['date_heure']}"
        
        if kind == 'resume':
            prompt_template = load_prompt('resume')
            prompt = render_prompt(
                prompt_template,
                matiere=m['matiere'],
                titre=m['titre'],
                date_heure=m['date_heure'],
                text=text
            )
            title = 'Résumé du cours'
        else:
            prompt_template = load_prompt('fiche')
            prompt = render_prompt(
                prompt_template,
                matiere=m['matiere'],
                titre=m['titre'],
                date_heure=m['date_heure'],
                text=text
            )
            title = 'Fiche de révision'
        
        self.bus.status.emit(f'Génération de : {title}…')
        
        client = MistralClient(k)
        try:
            content = client.chat(
                prompt,
                model=self.c.get('mistral_text_model', 'mistral-small-latest').strip() or 'mistral-small-latest',
                temperature=0.2
            )
            
            out = DOC / (p.stem + ('_resume.docx' if kind == 'resume' else '_fiche_revision.docx'))
            
            # Génération du document Word
            from docx import Document
            from docx.shared import Pt
            d = Document()
            d.add_heading(f"{title} - {m['matiere']}", 0)
            d.add_paragraph(m['titre'])
            d.add_paragraph(f"Date du cours : {m['date_heure']}")
            
            for line in content.splitlines():
                x = line.strip()
                if not x:
                    continue
                if x.startswith('### '):
                    d.add_heading(x[4:], 3)
                elif x.startswith('## '):
                    d.add_heading(x[3:], 2)
                elif x.startswith('# '):
                    d.add_heading(x[2:], 1)
                elif re.match(r'^[-*] ', x):
                    d.add_paragraph(x[2:], style='List Bullet')
                elif re.match(r'^\d+[.)] ', x):
                    d.add_paragraph(re.sub(r'^\d+[.)] ', '', x), style='List Number')
                else:
                    d.add_paragraph(x.replace('**', ''))
            
            d.styles['Normal'].font.name = 'Aptos'
            d.styles['Normal'].font.size = Pt(11)
            d.save(out)
            
            self.bus.done.emit(('doc', out))
        except Exception as e:
            self.bus.error.emit(str(e))

    def bg(self, f, *a):
        """Exécute une fonction en arrière-plan."""
        self.progress_bar.setVisible(True)
        self.progress_bar.setValue(0)
        
        def worker():
            try:
                f(*a)
            except Exception as e:
                self.bus.error.emit(str(e))
            finally:
                self.progress_bar.setVisible(False)
        
        threading.Thread(target=worker, args=(), daemon=True).start()

    def _finished(self, x):
        """Gère la fin d'une tâche."""
        k, p = x
        if k == 'tr':
            self.tr_page.fmtr.refresh()
        else:
            self.doc_page.fmdoc.refresh()
        self._refresh()
        self.statusBar().showMessage('Terminé : ' + p.name)

    def _err(self, s: str):
        """Affiche une erreur."""
        from mistral.exceptions import MistralAuthError, MistralQuotaError, NetworkError
        
        if isinstance(s, str):
            # Vérifier si le message correspond à une exception connue
            if "Clé API Mistral invalide" in s:
                QMessageBox.critical(
                    self, "Erreur d'authentification",
                    "Votre clé API Mistral est invalide ou a expiré.\n\n"
                    "Vérifiez-la dans : Paramètres → Clé API Mistral."
                )
            elif "Quota Mistral épuisé" in s:
                QMessageBox.critical(
                    self, "Quota épuisé",
                    "Votre quota Mistral est épuisé.\n\n"
                    "Connectez-vous à la [console Mistral](https://console.mistral.ai/) pour vérifier votre abonnement."
                )
            elif "Problème de connexion réseau" in s:
                QMessageBox.critical(
                    self, "Problème réseau",
                    "Impossible de se connecter à Mistral AI.\n\n"
                    "Vérifiez votre connexion Internet et réessayez."
                )
            else:
                QMessageBox.critical(self, "Erreur", s)
        else:
            QMessageBox.critical(self, "Erreur", str(s))

    def _refresh(self):
        """Rafraîchit les listes de fichiers."""
        if hasattr(self, 'rec_page'):
            self.rec_page._refresh()
        if hasattr(self, 'tr_page'):
            self.tr_page._refresh()
        if hasattr(self, 'doc_page'):
            self.doc_page._refresh()

    def showEvent(self, e):
        """Rafraîchit les listes au affichage."""
        super().showEvent(e)
        self._refresh()

    def closeEvent(self, e):
        """Arrête l'enregistrement à la fermeture."""
        self.r.stop()
        e.accept()


def main():
    """Point d'entrée de l'application."""
    app = QApplication(sys.argv)
    app.setStyleSheet(STYLE)
    window = MainWindow()
    window.show()
    sys.exit(app.exec())


if __name__ == '__main__':
    main()
